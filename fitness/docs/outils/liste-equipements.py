# -*- coding: utf-8 -*-
"""
Génère le tableur « Équipements et espaces bien-être » à envoyer aux équipes,
pré-rempli avec les données du lab.

    python3 docs/outils/liste-equipements.py        (depuis fitness/)

Rien n'est saisi ici : tout est lu dans data/data.js, et le fichier se
régénère à chaque fois que les données bougent.

Une ligne = un équipement, une croix par club qui le propose. C'est
exactement la forme du document client du 2026-09-17, en plus large : il ne
couvrait que le bien-être, celui-ci couvre toute la liste, parce que ce sont
les mêmes personnes qui savent, et qu'un deuxième document dans trois
semaines pour le parking serait un deuxième aller-retour.

La colonne « Bien-être » n'est pas décorative : c'est elle qui décide de ce
qui apparaît sur /bien-etre/espaces et sur le bloc bien-être des pages club.
La cocher sur un équipement l'y fait entrer.

Dépendances : openpyxl, et node pour lire data.js.
"""
import json, os, subprocess, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

ICI = os.path.dirname(os.path.abspath(__file__))
RACINE = os.path.dirname(os.path.dirname(ICI))
DATA = os.path.join(RACINE, 'data', 'data.js')
SORTIE = os.path.join(ICI, 'equipements-harmony.xlsx')

LECTEUR = """
global.window = {}; require(process.argv[1]);
var D = window.DATA;
console.log(JSON.stringify({
  clubs: D.clubs.map(function (c) {
    return { id: c.id, nom: c.nom, cat: c.categorie, canton: c.canton,
             equipements: c.equipements }; }),
  equipements: D.referentiels.equipements.map(function (e) {
    return { id: e.id, nom: e.nom, bienEtre: !!e.bienEtre }; })
}));
"""

try:
    brut = subprocess.check_output(['node', '-e', LECTEUR, DATA])
except (OSError, subprocess.CalledProcessError) as e:
    sys.exit('Lecture de data.js impossible : %s' % e)
d = json.loads(brut.decode('utf-8'))
clubs, equipements = d['clubs'], d['equipements']

CAT = {'gym': 'Gym', 'essential': 'Essential', 'premium': 'Premium'}
OUINON = ['Oui', 'Non']
COCHE = ['x']
STATUTS = ['Actif', 'Nouveau', 'À retirer']

# Les questions ouvertes sont posées dans le fichier, sur la ligne concernée :
# une question dans un e-mail se perd, une question dans la cellule se lit au
# moment où on remplit la ligne.
REMARQUES = {
    'jacuzzi': "Le document que vous nous avez transmis dit « Jacuzzi ou jet "
               "massant » en une seule ligne. Nous en faisons deux. Merci de "
               "préciser club par club, ou de nous dire s'il faut n'en faire "
               "qu'un seul équipement.",
    'jets':    "Même question que pour le jacuzzi : une ligne chez vous, deux "
               "chez nous.",
    'hammam':  "Veyrier : le document ne cochait pas le hammam, nous l'avons "
               "donc retiré. C'est la seule suppression de l'opération, et le "
               "document est ancien : à confirmer.",
}

POLICE = 'Arial'
NOIR = Font(name=POLICE, size=10)
BLANC = Font(name=POLICE, size=10, bold=True, color='FFFFFF')
TITRE = Font(name=POLICE, size=14, bold=True)
GRIS = Font(name=POLICE, size=10, color='767676')
H2 = Font(name=POLICE, size=11, bold=True, color='1F3864')
F_ENTETE = PatternFill('solid', fgColor='1F3864')
F_CLUB = PatternFill('solid', fgColor='2E5C8A')
F_CLE = PatternFill('solid', fgColor='5B5B5B')
F_ASAISIR = PatternFill('solid', fgColor='FFF7DC')
F_CALC = PatternFill('solid', fgColor='EFEFEF')
BORD = Border(*[Side(style='thin', color='D0D0D0')] * 4)

wb = Workbook()

# ------------------------------------------------------------------ Listes
ls = wb.create_sheet('Listes')
colonnes = [('OuiNon', OUINON), ('Coche', COCHE), ('Statuts', STATUTS)]
for i, (titre, vals) in enumerate(colonnes, start=1):
    c = ls.cell(row=1, column=i, value=titre); c.font = BLANC; c.fill = F_ENTETE
    for j, v in enumerate(vals, start=2):
        ls.cell(row=j, column=i, value=v).font = NOIR
    ls.column_dimensions[get_column_letter(i)].width = 24
ls['E1'] = "Ces listes alimentent les menus déroulants de l'onglet Équipements."
ls['E1'].font = GRIS

def plage(i, n):
    L = get_column_letter(i)
    return 'Listes!$%s$2:$%s$%d' % (L, L, n + 1)

# ------------------------------------------------------------------ Équipements
ws = wb.create_sheet('Équipements', 0)
entetes = ['Identifiant', 'Équipement', 'Statut', 'Bien-être']
NB_FIXES = len(entetes)
for c in clubs:
    entetes.append('%s\n(%s)' % (c['nom'], CAT[c['cat']]))
entetes += ['Nb clubs', 'Contrôle', 'Remarques']

G = NB_FIXES + 1                      # première colonne club
P = NB_FIXES + len(clubs)             # dernière colonne club
COL_NB, COL_CTRL, COL_REM = P + 1, P + 2, P + 3
LG, LP = get_column_letter(G), get_column_letter(P)
LN, LC = get_column_letter(COL_NB), get_column_letter(COL_CTRL)

for i, t in enumerate(entetes, start=1):
    c = ws.cell(row=1, column=i, value=t)
    c.font = BLANC
    c.fill = F_CLE if i == 1 else (F_CLUB if G <= i <= P else F_ENTETE)
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws.row_dimensions[1].height = 36

VIDES = 10
derniere = 1 + len(equipements) + VIDES

for k, e in enumerate(equipements):
    r = 2 + k
    ws.cell(row=r, column=1, value=e['id'])
    ws.cell(row=r, column=2, value=e['nom'])
    ws.cell(row=r, column=3, value='Actif')
    ws.cell(row=r, column=4, value='Oui' if e['bienEtre'] else 'Non')
    for j, c in enumerate(clubs):
        if e['id'] in c['equipements']:
            ws.cell(row=r, column=G + j, value='x')
    if e['id'] in REMARQUES:
        cel = ws.cell(row=r, column=COL_REM, value=REMARQUES[e['id']])
        cel.font = Font(name=POLICE, size=10, color='9C3B00')

for r in range(2, derniere + 1):
    neuve = r > 1 + len(equipements)
    for i in range(1, COL_REM + 1):
        c = ws.cell(row=r, column=i)
        if c.font is None or not c.font.color:
            c.font = GRIS if i == 1 else NOIR
        c.border = BORD
        c.alignment = Alignment(
            horizontal='center' if G <= i <= P else 'left',
            vertical='center', wrap_text=(i == COL_REM))
        if neuve and i != 1:
            c.fill = F_ASAISIR
    nb = ws.cell(row=r, column=COL_NB,
                 value='=IF(B{r}="","",COUNTA({a}{r}:{b}{r}))'.format(
                     r=r, a=LG, b=LP))
    nb.fill = F_CALC; nb.font = GRIS
    nb.alignment = Alignment(horizontal='center', vertical='center')
    ctrl = ws.cell(row=r, column=COL_CTRL, value=(
        '=IF(B{r}="","",'
        'IF(C{r}="","Statut à choisir",'
        'IF(D{r}="","Bien-être : oui ou non",'
        'IF({n}{r}=0,"Aucun club coché","OK"))))'
    ).format(r=r, n=LN))
    ctrl.fill = F_CALC; ctrl.font = GRIS
    ctrl.alignment = Alignment(horizontal='left', vertical='center')

def valide(formule, col_debut, col_fin=None):
    dv = DataValidation(type='list', formula1=formule, allow_blank=True,
                        showErrorMessage=True,
                        error='Choisissez une valeur dans la liste.',
                        errorTitle='Valeur hors liste')
    ws.add_data_validation(dv)
    a = get_column_letter(col_debut); b = get_column_letter(col_fin or col_debut)
    dv.add('%s2:%s%d' % (a, b, derniere))

valide(plage(3, len(STATUTS)), 3)
valide(plage(1, len(OUINON)), 4)
valide(plage(2, len(COCHE)), G, P)

rouge = PatternFill('solid', fgColor='FCE4E4')
bleu = PatternFill('solid', fgColor='E4EEF7')
vert = PatternFill('solid', fgColor='E8F3E8')
ws.conditional_formatting.add(
    '%s2:%s%d' % (LC, LC, derniere),
    FormulaRule(formula=['AND($%s2<>"",$%s2<>"OK")' % (LC, LC)], fill=rouge,
                font=Font(name=POLICE, size=10, bold=True, color='9C0006')))
for col in range(G, P + 1):
    L = get_column_letter(col)
    # La couleur d'une croix vient de la colonne Bien-être : on voit d'un coup
    # d'oeil ce qui alimente les pages Bien-être et ce qui reste un équipement.
    ws.conditional_formatting.add('%s2:%s%d' % (L, L, derniere), FormulaRule(
        formula=['AND(%s2<>"",$D2="Oui")' % L], fill=bleu,
        font=Font(name=POLICE, size=10, bold=True, color='1F3864')))
    ws.conditional_formatting.add('%s2:%s%d' % (L, L, derniere), FormulaRule(
        formula=['AND(%s2<>"",$D2="Non")' % L], fill=vert,
        font=Font(name=POLICE, size=10, color='2E5C2E')))

ws.freeze_panes = 'C2'
ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(COL_REM), derniere)
largeurs = {1: 20, 2: 26, 3: 12, 4: 12, COL_NB: 10, COL_CTRL: 26, COL_REM: 52}
for col in range(G, P + 1):
    largeurs[col] = 13
for col, w in largeurs.items():
    ws.column_dimensions[get_column_letter(col)].width = w

# ------------------------------------------------------------------ Textes
tx = wb.create_sheet('Textes')
for i, t in enumerate(['Identifiant', 'Équipement',
                       'Description (2 à 3 lignes, site)',
                       "Conditions d'accès", 'Règles d\'usage'], start=1):
    c = tx.cell(row=1, column=i, value=t)
    c.font = BLANC; c.fill = F_CLE if i == 1 else F_ENTETE
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
for k, e in enumerate(equipements):
    r = 2 + k
    tx.cell(row=r, column=1, value=e['id']).font = GRIS
    tx.cell(row=r, column=2, value=e['nom']).font = NOIR
    for i in range(3, 6):
        c = tx.cell(row=r, column=i); c.fill = F_ASAISIR; c.font = NOIR
        c.alignment = Alignment(wrap_text=True, vertical='top')
tx.freeze_panes = 'C2'
for col, w in {1: 20, 2: 26, 3: 70, 4: 44, 5: 44}.items():
    tx.column_dimensions[get_column_letter(col)].width = w
tx['G1'] = ("Les textes du site. Aujourd'hui remplacés par du faux texte sur "
            "la page Espaces wellness. Moins urgent que l'onglet Équipements.")
tx['G1'].font = GRIS
tx['G2'] = ("Conditions d'accès et règles d'usage : si elles changent d'un "
            "club à l'autre, dites-le ici et donnez le détail par club.")
tx['G2'].font = GRIS

# ------------------------------------------------------------------ Clubs
cl = wb.create_sheet('Clubs')
for i, t in enumerate(['Identifiant', 'Club', 'Catégorie', 'Canton',
                       'Nb équipements'], start=1):
    c = cl.cell(row=1, column=i, value=t)
    c.font = BLANC; c.fill = F_CLE if i == 1 else F_ENTETE
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
for k, c in enumerate(clubs):
    r = 2 + k
    cl.cell(row=r, column=1, value=c['id']).font = GRIS
    cl.cell(row=r, column=2, value=c['nom']).font = NOIR
    cl.cell(row=r, column=3, value=CAT[c['cat']]).font = NOIR
    cl.cell(row=r, column=4, value=c.get('canton') or '').font = NOIR
    L = get_column_letter(G + k)
    n = cl.cell(row=r, column=5, value="=COUNTA('Équipements'!%s2:%s%d)" % (
        L, L, derniere))
    n.fill = F_CALC; n.font = GRIS
    n.alignment = Alignment(horizontal='center')
for col, w in {1: 20, 2: 26, 3: 14, 4: 14, 5: 16}.items():
    cl.column_dimensions[get_column_letter(col)].width = w
cl['G1'] = ("Une colonne par club dans l'onglet Équipements, dans cet ordre. "
            "Le compte se met à jour tout seul.")
cl['G1'].font = GRIS

# ------------------------------------------------------------------ Mode d'emploi
me = wb.create_sheet("Mode d'emploi", 0)
me.column_dimensions['A'].width = 4
me.column_dimensions['B'].width = 118
contenu = [
    ('titre', 'Équipements et espaces bien-être'),
    ('gris', "Document de travail Harmony · pré-rempli avec ce que nous avons "
             "aujourd'hui, à vérifier et à compléter"),
    ('vide', ''),
    ('h2', "Ce qu'on vous demande"),
    ('p', "Onglet « Équipements » : une ligne par équipement, et une croix "
          "dans la colonne de chaque club qui le propose. C'est la seule "
          "chose urgente. Les onglets « Textes » et « Clubs » peuvent "
          "attendre."),
    ('p', "La liste des équipements est fermée, et c'est voulu : sinon un "
          "club écrit « Sauna », un autre « sauna finlandais », et les "
          "filtres du site ne marchent plus. S'il manque un équipement, "
          "ajoutez une ligne en bas, dans les lignes jaunes, et laissez-nous "
          "l'identifiant vide : on s'en occupe."),
    ('vide', ''),
    ('h2', "La colonne qui décide de tout : Bien-être"),
    ('p', "Oui ou Non. Les équipements marqués Oui, et eux seuls, "
          "apparaissent sur la page « Espaces wellness » du site et dans le "
          "bloc bien-être de chaque page club. Les autres restent de simples "
          "équipements, listés sur la page du club."),
    ('p', "Aujourd'hui cinq sont à Oui : piscine, sauna, hammam, jacuzzi, "
          "jets massants. La piscine est la plus discutable : elle sert aussi "
          "aux cours aquatiques et à l'école de natation, et l'annoncer trois "
          "fois dans trois discours différents ne rend service à personne. "
          "Votre avis nous intéresse sur cette ligne en particulier."),
    ('vide', ''),
    ('h2', "Deux questions posées dans le fichier"),
    ('p', "Elles sont écrites en orange dans la colonne « Remarques », sur la "
          "ligne concernée. Jacuzzi et jets massants : votre document n'en "
          "faisait qu'une ligne, « Jacuzzi ou jet massant », nous en faisons "
          "deux. Hammam à Veyrier : nous l'avons retiré parce que le document "
          "ne le cochait pas, et ce document est ancien."),
    ('vide', ''),
    ('h2', 'Les autres colonnes'),
    ('p', "Identifiant · notre clé technique. Elle ne change jamais, même si "
          "l'équipement est renommé. Ne rien y écrire, ne pas la supprimer : "
          "c'est elle qui permet de recharger ce fichier sans tout recasser."),
    ('p', "Équipement · le nom exact, tel qu'il doit apparaître sur le site. "
          "C'est le seul endroit où ce nom se décide."),
    ('p', "Statut · Actif, Nouveau, ou À retirer. Ne supprimez pas une "
          "ligne : marquez-la À retirer, on saura quoi faire de l'existant."),
    ('p', "Une colonne par club · une croix si le club le propose, rien "
          "sinon. Une case vide veut dire « non », pas « je ne sais pas » : "
          "s'il y a un doute, dites-le en Remarques."),
    ('p', "Nb clubs et Contrôle · calculés, ne rien y écrire. Le contrôle "
          "devient rouge quand il manque quelque chose."),
    ('vide', ''),
    ('h2', "Le point le plus facile à rater"),
    ('p', "Un équipement hors service depuis six mois est toujours un "
          "équipement du club, et le site l'annoncera. Si un sauna est fermé "
          "pour travaux, ne décochez pas la case : dites-le en Remarques. "
          "Décocher revient à dire au visiteur qu'il n'y en a pas, et une "
          "absence annoncée à tort est pire qu'un silence."),
    ('vide', ''),
    ('h2', 'Les couleurs'),
    ('p', "Jaune pâle · à remplir. Bleu · une croix sur un équipement "
          "bien-être. Vert · une croix sur un équipement ordinaire. Rouge · "
          "une incohérence à regarder. Orange · une question qu'on vous pose."),
    ('vide', ''),
    ('gris', "Les lignes déjà remplies viennent du document de disponibilité "
             "par club que vous nous avez transmis, et de ce que nous avions "
             "avant lui. Elles sont à vérifier, pas à prendre pour argent "
             "comptant."),
]
r = 2
for genre, txt in contenu:
    c = me.cell(row=r, column=2, value=txt)
    c.font = {'titre': TITRE, 'h2': H2, 'gris': GRIS}.get(genre, NOIR)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    if genre not in ('p', 'gris'):
        me.row_dimensions[r].height = 20
    r += 1

del wb['Sheet']
wb.save(SORTIE)
print('écrit : %s (%d équipements, %d clubs, %d lignes vides)'
      % (SORTIE, len(equipements), len(clubs), VIDES))
